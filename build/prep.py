"""Consolidate every contact-bearing sheet of IFA_Meeting_List.xlsx into one master list."""
import json, datetime, re
import openpyxl

SRC = 'source.xlsx'
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

def s(v):
    if v is None:
        return ''
    if isinstance(v, str):
        return v.replace('\xa0', ' ').strip()
    return v

def key(name, email=''):
    """People are keyed on their name with spacing/case ignored, so 'Yeo Chinwei' and
    'Yeo Chin Wei' collapse into one contact across the five source sheets."""
    return re.sub(r'[^a-z]', '', s(name).lower())

# ---------------------------------------------------------------- firm mapping
FIRM_MAP = {
    'synergy': 'Synergy', 'synergy fa': 'Synergy',
    'financial alliance': 'Financial Alliance',
    'phillip securities': 'Phillip Capital', 'phillip capital': 'Phillip Capital',
    'infinity fa': 'Infinity FA',
    'finexis advisory': 'Finexis Advisory', 'finexis': 'Finexis Advisory',
    'pias': 'PIAS',
    'ifast': 'iFast', 'ifast (product)': 'iFast', 'ifastgm': 'iFast',
    'singlife': 'Singlife', 'singlife financial advisers': 'Singlife',
    'alpha wealth': 'Alpha Wealth',
    'gyc financial advisory': 'GYC Financial Advisory',
    'sg alliance': 'SG Alliance',
    'ipp fa': 'IPP FA',
    'hsbc life': 'HSBC Life',
    'h2o am': 'H2O AM (internal)',
    'client of norman ngai': 'Other / Client',
    'neba private': 'NEBA Private',
    'promiseland': 'PromiseLand', 'promiseland financial advisers': 'PromiseLand',
    'promiseland financial advisory': 'PromiseLand',
    'great eastern life': 'Great Eastern Life',
    'manulife fa': 'Manulife FA', 'manulife financial advisers': 'Manulife FA',
    'aia financial adviser': 'AIA FA',
    'income advisory financial advisers': 'Income Advisory',
}
def firm(v):
    t = s(v)
    if not t:
        return 'Unknown'
    return FIRM_MAP.get(t.lower(), t)

# --------------------------------------------------------------- record shape
FIELDS = ['name', 'email', 'firm', 'ctype', 'source', 'met', 'lastmtg', 'mtgs',
          'invested', 'aum', 'mailsent', 'linkedin', 'e1', 'e2', 'e3', 'e4', 'e5', 'notes']
NOT_INV = 'Not Invited'

master, order = {}, []

def add(rec):
    k = key(rec['name'])
    # same name at a different firm = a different person, so give them their own key
    while k in master and rec.get('firm') and master[k].get('firm') \
            and master[k]['firm'] not in ('Unknown', rec['firm']):
        k += '#'
    if k in master:
        cur = master[k]
        for f in FIELDS:
            if not cur.get(f) and rec.get(f):
                cur[f] = rec[f]
        for f in ('e1', 'e2', 'e3', 'e4', 'e5'):
            if cur.get(f, NOT_INV) == NOT_INV and rec.get(f, NOT_INV) != NOT_INV:
                cur[f] = rec[f]
        if rec.get('notes') and rec['notes'] not in (cur.get('notes') or ''):
            cur['notes'] = '; '.join(x for x in [cur.get('notes'), rec['notes']] if x)
        if rec.get('source') and rec['source'] not in cur['source']:
            cur['source'] += ' + ' + rec['source']
        return cur
    base = {f: '' for f in FIELDS}
    base.update({'e1': NOT_INV, 'e2': NOT_INV, 'e3': NOT_INV, 'e4': NOT_INV, 'e5': NOT_INV,
                 'mtgs': 0, 'aum': ''})
    base.update(rec)
    master[k] = base
    order.append(k)
    return base

# ------------------------------------------------- 1. Attendance = May-26 event
attend = {}
for r in wb['Attendance'].iter_rows(min_row=2, values_only=True):
    if not s(r[0]):
        continue
    attend[key(r[0])] = (s(r[3]) == 1 or s(r[3]) == '1')

# --------------------------------------------------------- 2. Meeting List core
# Oct-25 event  : col D "Events"  (Yes / No / Not Available / blank)
# May-26 event  : col J "Invited", col K "Accepted", + Attendance sheet
E1_MAP = {'Yes': 'Attended', 'No': 'No Reply', 'Not Available': 'Declined', '': NOT_INV}

for r in wb['Meeting List'].iter_rows(min_row=3, values_only=True):
    name = s(r[0])
    if not name:
        continue
    email = s(r[1])
    if str(email).lower() in ('n/a', 'na'):
        email = ''
    notes = []

    e1 = E1_MAP.get(s(r[3]), NOT_INV)

    invited_raw, accepted_raw = s(r[9]), s(r[10])
    e2, e3 = NOT_INV, NOT_INV
    k = key(name)
    if k in attend:
        e2 = 'Attended' if attend[k] else 'No Show'
    elif accepted_raw in ('Declined',):
        e2 = 'Declined'
    elif accepted_raw == 'Not Available':
        e2 = 'Declined'
        notes.append('May-26: not available')
    elif accepted_raw == 'Bounced back':
        e2 = 'No Reply'
        notes.append('May-26: email bounced back')
    elif invited_raw in (1, '1') or s(r[11]) in (1, '1'):
        e2 = 'No Reply'

    if invited_raw == 'Macro Lunch':
        e3 = 'Invited'
        e2 = NOT_INV if e2 == 'No Reply' else e2
    elif invited_raw == 'Overseas':
        e2 = 'Declined'
        notes.append('Overseas')
    elif invited_raw == 'TBC':
        e2 = 'Invited'

    met_raw = s(r[5])
    met = {'Yes': 'Yes', 'yes': 'Yes', 'No': 'No', 'TBC': 'TBC', '': ''}.get(met_raw, met_raw)
    last = r[6] if isinstance(r[6], datetime.datetime) else ''
    inv_raw = s(r[7])
    invested = {'Yes': 'Yes', 'No': 'No', '?': 'TBC', '': ''}.get(inv_raw, inv_raw)

    add({'name': name, 'email': email, 'firm': firm(r[2]), 'ctype': 'Adviser',
         'source': 'Meeting List', 'met': met, 'lastmtg': last,
         'mtgs': 1 if met == 'Yes' else 0, 'invested': invested,
         'e1': e1, 'e2': e2, 'e3': e3, 'notes': '; '.join(notes)})

# --------------------------------------------------------------- 3. ILP => AUM
for r in wb['ILP Data'].iter_rows(min_row=2, values_only=True):
    name = s(r[0])
    if not name or name.lower().startswith('total'):
        continue
    aum = r[2] if isinstance(r[2], (int, float)) else ''
    note = '' if isinstance(r[2], (int, float)) else (f'ILP: {s(r[2])}' if s(r[2]) else '')
    add({'name': name, 'email': '', 'firm': firm(r[1]), 'ctype': 'Adviser',
         'source': 'ILP Data', 'aum': aum, 'notes': note})

# ------------------------------------------------------- 4. Mass Email to PIAS
for r in wb['Mass Email PIAS'].iter_rows(min_row=2, values_only=True):
    name = s(r[0])
    if not name:
        continue
    replied = s(r[5]) == 'Yes'
    mtg_date = r[6] if isinstance(r[6], datetime.datetime) else ''
    add({'name': name, 'email': s(r[2]), 'firm': firm(r[1]), 'ctype': 'Adviser',
         'source': 'Mass Email PIAS',
         'mailsent': 'Yes' if s(r[3]).startswith('Yes') else 'No',
         'met': 'Yes' if mtg_date else '', 'lastmtg': mtg_date, 'mtgs': 1 if mtg_date else 0,
         'notes': 'Replied to mass email' if replied else ''})

# ------------------------------------------------------------ 5. Branch Directors
LI_MAP = {'Yes': 'Yes', 'No': 'No', "Can't find LinkedIn": 'No LinkedIn found', '': ''}
cur_firm = ''
for r in wb['Branch Directors'].iter_rows(min_row=2, values_only=True):
    if s(r[0]):
        cur_firm = firm(r[0])
    if s(r[0]) == 'deg':
        continue
    branch, director, consultant = s(r[1]), s(r[2]), s(r[3])
    li_contacted = LI_MAP.get(s(r[4]), s(r[4]))
    li_msg = s(r[6])
    if li_msg == 'already in CRM' and li_contacted in ('', 'No'):
        li_contacted = 'Yes'
    elif li_msg == 'Yes' and li_contacted in ('', 'No'):
        li_contacted = 'Yes'
    note = ' | '.join(x for x in [f'Branch: {branch}' if branch else '',
                                  s(r[5]), 'Already in CRM' if li_msg == 'already in CRM' else '',
                                  s(r[8])] if x)
    if director:
        add({'name': director, 'email': s(r[7]), 'firm': cur_firm or 'Unknown',
             'ctype': 'Branch Director', 'source': 'Branch Directors',
             'linkedin': li_contacted, 'notes': note})
    if consultant:
        add({'name': consultant, 'email': '', 'firm': cur_firm or 'Unknown',
             'ctype': 'Consultant', 'source': 'Branch Directors',
             'notes': (f'Reports to {director}' if director else '') +
                      (f' | Branch: {branch}' if branch else '')})

# --------------------------------------------------------------------- tidy up
rows = [master[k] for k in order]
for rec in rows:
    # a last-meeting date can only exist if the meeting happened: 9 rows of the source
    # carried a date while 'Meeting (Thomas)' still said No
    if rec['lastmtg'] and rec['met'] != 'Yes':
        rec['met'] = 'Yes'
    if rec['mtgs'] and rec['met'] != 'Yes':
        rec['met'] = 'Yes'
    if not rec['met']:
        rec['met'] = 'Yes' if rec['mtgs'] else 'No'
    if rec['met'] == 'Yes' and not rec['mtgs']:
        rec['mtgs'] = 1
    if not rec['invested']:
        rec['invested'] = 'No'
    if not rec['mailsent']:
        rec['mailsent'] = 'Yes' if rec['source'].startswith('Mass Email') else ''
    if isinstance(rec['lastmtg'], datetime.datetime):
        rec['lastmtg'] = rec['lastmtg'].date().isoformat()

rows.sort(key=lambda r: (r['firm'].lower(), r['name'].lower()))

from collections import Counter
print('total contacts:', len(rows))
print('firms:', len(set(r['firm'] for r in rows)))
print(Counter(r['firm'] for r in rows).most_common())
print('met:', Counter(r['met'] for r in rows))
print('E1:', Counter(r['e1'] for r in rows))
print('E2:', Counter(r['e2'] for r in rows))
print('E3:', Counter(r['e3'] for r in rows))
print('types:', Counter(r['ctype'] for r in rows))
print('sources:', Counter(r['source'] for r in rows))
print('with AUM:', sum(1 for r in rows if r['aum']), 'total AUM', sum(r['aum'] for r in rows if r['aum']))
print('total meetings:', sum(r['mtgs'] for r in rows))

json.dump(rows, open('contacts.json', 'w'), indent=0)
